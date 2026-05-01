import streamlit as st
import os
import random
import io
import time
from dotenv import load_dotenv
from docx import Document

# Imports LangChain & Groq
from langchain_groq import ChatGroq
from langchain_community.vectorstores import FAISS
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnablePassthrough
from langchain_core.output_parsers import StrOutputParser

# Import de ton processeur personnalisé
try:
    from core.processor import DocumentProcessor
except ImportError:
    st.error("Le module core.processor est introuvable.")

# --- CONFIGURATION DES CLES ---
load_dotenv()
groq_api_key = os.getenv("GROQ_API_KEY")
google_api_key = os.getenv("GOOGLE_API_KEY") 

st.set_page_config(page_title="Nexus AI - Groq Edition", page_icon="⚡", layout="wide")

# --- DONNÉES PAYS ---
COUNTRIES = {
    "Tunisie (+216)": "+216", "France (+33)": "+33", "Maroc (+212)": "+212",
    "Algérie (+213)": "+213", "Canada (+1)": "+1", "Belgique (+32)": "+32"
}

# --- INITIALISATION DU SESSION STATE ---
if "auth_status" not in st.session_state: st.session_state.auth_status = "start" # On commence par la page d'accueil
if "processor" not in st.session_state: st.session_state.processor = DocumentProcessor()
if "vector_db" not in st.session_state: st.session_state.vector_db = None
if "messages" not in st.session_state: st.session_state.messages = []
if "user_name" not in st.session_state: st.session_state.user_name = "Utilisateur"
if "user_email" not in st.session_state: st.session_state.user_email = ""

# --- FONCTION EXPORT WORD ---
def generate_word_doc(text):
    doc = Document()
    doc.add_heading('Nexus AI - Rapport d\'Analyse', 0)
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- STYLE CSS ---
st.markdown("""
    <style>
    header {visibility: hidden;}
    .profile-container {
        position: fixed; top: 10px; right: 20px; z-index: 10000;
        display: flex; align-items: center; gap: 10px;
        background-color: #1e293b; padding: 5px 15px;
        border-radius: 30px; border: 1px solid #334155; color: white;
    }
    .profile-pic {
        width: 35px; height: 35px; background-color: #f59e0b;
        color: black; border-radius: 50%; display: flex;
        justify-content: center; align-items: center; font-weight: bold;
    }
    .main-title { text-align: center; margin-top: 50px; font-weight: 800; color: #f59e0b; font-size: 4rem; }
    .sub-title { text-align: center; color: #94a3b8; font-size: 1.5rem; margin-bottom: 30px; }
    [data-testid="stSidebar"] { background-color: #0f172a; }
    </style>
    """, unsafe_allow_html=True)

# --- 1. LOGIQUE D'AUTHENTIFICATION ---

# ÉTAPE A : PAGE D'ACCUEIL (Titre + Bouton Connexion)
if st.session_state.auth_status == "start":
    st.markdown("<h1 class='main-title'>Nexus AI Intelligence</h1>", unsafe_allow_html=True)
    st.markdown("<p class='sub-title'>Analyse de documents et intelligence augmentée</p>", unsafe_allow_html=True)
    
    _, col, _ = st.columns([1, 0.6, 1])
    with col:
        if st.button("🚀 Se connecter à Nexus", use_container_width=True):
            st.session_state.auth_status = "login"
            st.rerun()

# ÉTAPE B : FORMULAIRE D'INSCRIPTION
elif st.session_state.auth_status == "login":
    _, col, _ = st.columns([1, 1.2, 1])
    with col:
        st.markdown("<h2 style='text-align: center;'>✨ Inscription Nexus</h2>", unsafe_allow_html=True)
        with st.form("inscription"):
            name = st.text_input("Nom complet")
            email = st.text_input("Email")
            c1, c2 = st.columns([0.4, 0.6])
            with c1: country = st.selectbox("Pays", list(COUNTRIES.keys()))
            with c2: phone = st.text_input("Téléphone")
            
            if st.form_submit_button("Continuer", use_container_width=True):
                if name and email:
                    st.session_state.user_name = name.capitalize()
                    st.session_state.user_email = email
                    st.session_state.generated_code = str(random.randint(1000, 9999))
                    st.session_state.auth_status = "verify"
                    st.rerun()

# ÉTAPE C : VÉRIFICATION CODE
elif st.session_state.auth_status == "verify":
    _, col, _ = st.columns([1, 1.2, 1])
    with col:
        st.info(f"💡 Code de test : {st.session_state.generated_code}")
        input_code = st.text_input("Code secret", type="password")
        if st.button("Valider"):
            if input_code == st.session_state.generated_code:
                st.session_state.auth_status = "authenticated"
                st.rerun()

# --- 2. INTERFACE PRINCIPALE (Après connexion) ---
elif st.session_state.auth_status == "authenticated":
    
    # Widget Profil Haut-Droit
    st.markdown(f"""
        <div class="profile-container">
            <span style="font-size: 14px;">{st.session_state.user_name}</span>
            <div class="profile-pic">{st.session_state.user_name[0]}</div>
        </div>
        """, unsafe_allow_html=True)

    # Sidebar
    with st.sidebar:
        st.markdown(f"### 📂 Dashboard")
        st.caption(f"Connecté en tant que : {st.session_state.user_email}")
        if st.button("➕ Nouvelle Discussion", use_container_width=True):
            st.session_state.messages = []
            st.session_state.vector_db = None
            st.rerun()
        st.divider()
        if st.button("🚪 Déconnexion"):
            st.session_state.auth_status = "start" # Retour à la page d'accueil
            st.rerun()

    st.markdown("<h1 style='text-align: center; color: #f59e0b; font-weight: 800; font-size: 3rem; margin-top: -30px;'>Nexus AI Intelligence</h1>", unsafe_allow_html=True)

    # Zone d'analyse
    _, mid_col, _ = st.columns([1, 2.5, 1])
    with mid_col:
        with st.expander("📥 Charger une source", expanded=(st.session_state.vector_db is None)):
            c1, c2 = st.columns(2)
            with c1: file = st.file_uploader("Document", type=["pdf", "png", "jpg"])
            with c2: url = st.text_input("YouTube URL")
            
            if st.button("Lancer l'analyse", use_container_width=True):
                with st.spinner("Analyse Nexus en cours..."):
                    try:
                        chunks = None
                        if file:
                            t_path = f"temp_{file.name}"
                            with open(t_path, "wb") as f: f.write(file.getbuffer())
                            chunks = st.session_state.processor.process_file(t_path)
                            os.remove(t_path)
                        elif url:
                            chunks = st.session_state.processor.process_file(url)
                        
                        if chunks:
                            embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
                            st.session_state.vector_db = FAISS.from_documents(chunks, embeddings)
                            st.success("Analyse terminée ! Vous pouvez poser vos questions.")
                    except Exception as e:
                        st.error(f"Erreur : {str(e)}")

    st.divider()

    # Zone de Chat
    _, chat_col, _ = st.columns([1, 2.5, 1])
    with chat_col:
        for i, m in enumerate(st.session_state.messages):
            with st.chat_message(m["role"]):
                st.markdown(m["content"])
                if m["role"] == "assistant":
                    st.download_button(
                        label="📄 Télécharger en Word",
                        data=generate_word_doc(m["content"]),
                        file_name=f"analyse_nexus_{i}.docx",
                        key=f"dl_{i}"
                    )

        if prompt := st.chat_input("Posez votre question..."):
            st.session_state.messages.append({"role": "user", "content": prompt})
            with st.chat_message("user"): st.markdown(prompt)

            with st.chat_message("assistant"):
                if st.session_state.vector_db:
                    with st.spinner("Nexus réfléchit..."):
                        llm = ChatGroq(temperature=0, model_name="llama-3.1-8b-instant", groq_api_key=groq_api_key)
                        retriever = st.session_state.vector_db.as_retriever()
                        prompt_t = ChatPromptTemplate.from_template("Contexte: {context}\n\nQuestion: {question}")
                        
                        chain = ({"context": retriever | (lambda docs: "\n\n".join(d.page_content for d in docs)), 
                                 "question": RunnablePassthrough()} | prompt_t | llm | StrOutputParser())
                        
                        response = chain.invoke(prompt)
                        st.markdown(response)
                        
                        st.download_button(
                            label="📄 Télécharger en Word",
                            data=generate_word_doc(response),
                            file_name="reponse_nexus.docx",
                            key=f"dl_new"
                        )
                        st.session_state.messages.append({"role": "assistant", "content": response})
                else:
                    st.warning("⚠️ Veuillez charger une source avant de poser une question.")
            st.rerun()