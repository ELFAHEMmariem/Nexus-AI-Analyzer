from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
from datetime import datetime

def generate_word_report(answer, query):
    """Génère un fichier Word professionnel à partir de la réponse de l'IA"""
    doc = Document()
    
    # --- Style du titre ---
    title = doc.add_heading('Rapport d\'Analyse Expert IA', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # --- Infos de session ---
    doc.add_paragraph(f"Date de génération : {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph("---------------------------------------------------------")
    
    # --- Section Question ---
    doc.add_heading('Question posée :', level=1)
    q_para = doc.add_paragraph()
    q_run = q_para.add_run(query)
    q_run.italic = True
    
    # --- Section Réponse ---
    doc.add_heading('Analyse de l\'Expert :', level=1)
    doc.add_paragraph(answer)
    
    # --- Pied de page ---
    doc.add_paragraph("\n\nDocument généré automatiquement par Expert Multimodal RAG.")
    
    # Sauvegarde dans un flux de bits (mémoire vive) pour Streamlit
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()